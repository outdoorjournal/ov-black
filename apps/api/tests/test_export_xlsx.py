"""XLSX rendering: sheet structure, cell values, formats, freeze panes."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from app.services.export.xlsx import build_xlsx
from openpyxl import load_workbook

from tests._export_fixtures import make_day, make_item, make_trip


def _trip() -> object:
    day1 = make_day(
        date(2026, 6, 1),
        (
            make_item(
                title="Fushimi Inari",
                type="experience",
                start="2026-06-01T09:00:00+09:00",
                duration_minutes=120,
                location="Kyoto, Japan",
                cost="100.00",
                currency="USD",
                cost_kind="total",
            ),
            make_item(
                title="Send-off dinner",
                type="meal",
                start="2026-06-01T19:00:00+09:00",
                duration_minutes=90,
                cost="50.00",
                currency="USD",
                cost_kind="per_person",
                notes=("Window table",),
            ),
        ),
        label="Day 1",
        night_hotel="Hotel Okura",
        subtotals={"USD": Decimal("250.00")},
    )
    collection = (
        make_item(
            title="Maybe: tea auction",
            start=None,
            duration_minutes=None,
            cost="200.00",
            currency="USD",
            cost_kind="total",
        ),
    )
    return make_trip(
        (day1,), collection=collection, totals={"USD": Decimal("450.00")}, party_size=3
    )


def test_build_xlsx_structure_and_values() -> None:
    data = build_xlsx(_trip())  # type: ignore[arg-type]
    assert data[:2] == b"PK"  # xlsx is a zip
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["Itinerary", "Summary", "Collection"]

    itinerary = wb["Itinerary"]
    assert itinerary.freeze_panes == "A2"
    header = [c.value for c in itinerary[1]]
    assert header[:8] == [
        "Day",
        "Date",
        "Start",
        "End",
        "Duration (min)",
        "Type",
        "Title",
        "Location",
    ]

    # First data row = the temple.
    row2 = {c.column_letter: c.value for c in itinerary[2]}
    assert row2["A"] == "Day 1"
    # openpyxl round-trips a python date as a datetime cell value.
    assert row2["B"] == datetime(2026, 6, 1)
    assert row2["C"] == "09:00"
    assert row2["D"] == "11:00"
    assert row2["G"] == "Fushimi Inari"
    assert row2["H"] == "Kyoto, Japan"
    assert row2["I"] == 100.0
    assert itinerary["B2"].number_format == "yyyy-mm-dd"
    assert itinerary["I2"].number_format == "#,##0.00"

    # The dinner note lands in Notes.
    row3 = {c.column_letter: c.value for c in itinerary[3]}
    assert row3["G"] == "Send-off dinner"
    assert row3["K"] == "per person"
    assert row3["M"] == "Window table"

    # Subtotal row.
    subtotal_titles = [
        row[6].value
        for row in itinerary.iter_rows()
        if row[6].value and "subtotal" in str(row[6].value)
    ]
    assert subtotal_titles == ["Day 1 subtotal"]


def test_summary_and_collection_sheets() -> None:
    data = build_xlsx(_trip())  # type: ignore[arg-type]
    wb = load_workbook(BytesIO(data))

    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows()}
    assert summary["Trip"] == "Kyoto in June"
    assert summary["Party size"] == 3
    assert summary["Total (USD)"] == 450.0
    assert summary["Version"] == "Official itinerary"

    collection = wb["Collection"]
    assert [c.value for c in collection[1]][:2] == ["Type", "Title"]
    assert collection["B2"].value == "Maybe: tea auction"
    assert collection["D2"].value == 200.0
