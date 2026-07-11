"""XLSX rendering of an :class:`~app.services.export.grouping.ExportTrip`.

A journal-like structured timeline: one row per item with times, locations,
and prices in aligned columns, day-subtotal rows, a Summary sheet with trip
metadata and per-currency totals, and a Collection sheet for unscheduled
ideas. Times are LOCAL wall-clock strings (each node's own zone offset is
already applied by the grouping layer) — Excel has no timezone type.
"""

from __future__ import annotations

from datetime import timedelta
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.export.grouping import ExportItem, ExportTrip

_HEADER_FONT = Font(bold=True)
_SUBTOTAL_FILL = PatternFill(start_color="FDEBDD", end_color="FDEBDD", fill_type="solid")
_CURRENCY_FORMAT = "#,##0.00"
_DATE_FORMAT = "yyyy-mm-dd"

_ITINERARY_HEADERS = (
    "Day",
    "Date",
    "Start",
    "End",
    "Duration (min)",
    "Type",
    "Title",
    "Location",
    "Cost",
    "Currency",
    "Cost kind",
    "Status",
    "Notes",
)
_ITINERARY_WIDTHS = (10, 12, 8, 8, 14, 12, 44, 30, 12, 10, 12, 12, 50)

_COLLECTION_HEADERS = (
    "Type",
    "Title",
    "Location",
    "Cost",
    "Currency",
    "Cost kind",
    "Status",
    "Notes",
)
_COLLECTION_WIDTHS = (12, 44, 30, 12, 10, 12, 12, 50)


def _write_header(sheet: Worksheet, headers: tuple[str, ...], widths: tuple[int, ...]) -> None:
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = _HEADER_FONT
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"


def _item_notes(item: ExportItem) -> str:
    parts: list[str] = []
    if item.journey is not None:
        parts.append(
            f"Day {item.journey.index} of {item.journey.total}"
            f" — part of {item.journey.parent_title}"
        )
    parts.extend(item.notes)
    if item.description:
        parts.append(item.description[:300])
    return " · ".join(parts)


def _cost_cells(item: ExportItem) -> tuple[Any, str | None, str | None]:
    if item.cost_amount is None or item.cost_currency is None:
        return (None, None, None)
    kind = "per person" if item.cost_kind == "per_person" else "total"
    return (float(item.cost_amount), item.cost_currency, kind)


def _end_time(item: ExportItem) -> str | None:
    if item.start_local is None or item.duration_minutes is None:
        return None
    return (item.start_local + timedelta(minutes=item.duration_minutes)).strftime("%H:%M")


def build_xlsx(trip: ExportTrip) -> bytes:
    """Render the whole trip into an .xlsx byte string."""
    workbook = Workbook()

    # ── Itinerary: the journal timeline ──────────────────────────────────
    sheet = workbook.active
    assert sheet is not None  # openpyxl always creates one
    sheet.title = "Itinerary"
    _write_header(sheet, _ITINERARY_HEADERS, _ITINERARY_WIDTHS)

    for day in trip.days:
        for item in day.items:
            cost, currency, kind = _cost_cells(item)
            sheet.append(
                [
                    day.label,
                    day.day_date,
                    item.start_local.strftime("%H:%M") if item.start_local else None,
                    _end_time(item),
                    item.duration_minutes,
                    item.type,
                    item.title,
                    item.location,
                    cost,
                    currency,
                    kind,
                    item.status,
                    _item_notes(item) or None,
                ]
            )
            row = sheet.max_row
            sheet.cell(row=row, column=2).number_format = _DATE_FORMAT
            if cost is not None:
                sheet.cell(row=row, column=9).number_format = _CURRENCY_FORMAT
        if day.subtotals:
            for currency_code, amount in sorted(day.subtotals.items()):
                sheet.append(
                    [day.label, day.day_date, None, None, None, None, f"{day.label} subtotal", None]
                    + [float(amount), currency_code, None, None, None]
                )
                row = sheet.max_row
                sheet.cell(row=row, column=2).number_format = _DATE_FORMAT
                sheet.cell(row=row, column=9).number_format = _CURRENCY_FORMAT
                for cell in sheet[row]:
                    cell.fill = _SUBTOTAL_FILL
                sheet.cell(row=row, column=7).font = _HEADER_FONT

    # ── Summary: trip metadata + totals ──────────────────────────────────
    summary = workbook.create_sheet("Summary")
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 60

    def _kv(key: str, value: Any) -> None:
        summary.append([key, value])
        summary.cell(row=summary.max_row, column=1).font = _HEADER_FONT
        summary.cell(row=summary.max_row, column=2).alignment = Alignment(wrap_text=True)

    _kv("Trip", trip.title)
    if trip.brief:
        _kv("Brief", trip.brief)
    _kv("Version", "Working copy (fork)" if trip.is_fork else "Official itinerary")
    if trip.timing_kind:
        _kv("Timing", trip.timing_kind)
    if trip.date_start:
        _kv("Starts", trip.date_start.isoformat())
    if trip.date_end:
        _kv("Ends", trip.date_end.isoformat())
    if trip.duration_nights:
        _kv("Nights", trip.duration_nights)
    if trip.timing_note:
        _kv("Timing note", trip.timing_note)
    _kv("Party size", trip.party_size)
    _kv("Days planned", len(trip.days))
    for currency_code, amount in sorted(trip.totals.items()):
        summary.append([f"Total ({currency_code})", float(amount)])
        summary.cell(row=summary.max_row, column=1).font = _HEADER_FONT
        summary.cell(row=summary.max_row, column=2).number_format = _CURRENCY_FORMAT
    if trip.totals:
        _kv("Pricing note", f"Per-person prices expanded × party of {trip.party_size}.")

    # ── Collection: unscheduled ideas ────────────────────────────────────
    collection = workbook.create_sheet("Collection")
    _write_header(collection, _COLLECTION_HEADERS, _COLLECTION_WIDTHS)
    for item in trip.collection:
        cost, currency, kind = _cost_cells(item)
        collection.append(
            [
                item.type,
                item.title,
                item.location,
                cost,
                currency,
                kind,
                item.status,
                _item_notes(item) or None,
            ]
        )
        if cost is not None:
            collection.cell(row=collection.max_row, column=4).number_format = _CURRENCY_FORMAT

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
